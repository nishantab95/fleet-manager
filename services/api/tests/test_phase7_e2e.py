from __future__ import annotations

from collections.abc import Iterator
from datetime import UTC, datetime, timedelta
from io import BytesIO
from typing import cast
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook  # type: ignore[import-untyped]
from sqlalchemy.orm import Session

from fleet_api.api.dependencies import get_object_storage, get_otp_provider
from fleet_api.auth.providers import FakeOtpProvider
from fleet_api.core.config import Settings
from fleet_api.db.models import CompanyMembership, User
from fleet_api.db.session import get_db as session_get_db
from fleet_api.main import create_app

pytestmark = pytest.mark.postgres


class InMemoryStorage:
    def __init__(self) -> None:
        self.objects: dict[str, bytes] = {}

    def put_private(self, *, object_key: str, content: bytes, content_type: str) -> str:
        del content_type
        self.objects[object_key] = content
        return object_key

    def delete_private(self, *, object_key: str) -> None:
        self.objects.pop(object_key, None)

    def read_private(self, *, object_key: str) -> tuple[bytes, str]:
        return self.objects[object_key], "image/jpeg"


def _settings() -> Settings:
    return Settings(
        environment="test",
        phone_default_region="IN",
        jwt_signing_key="test-signing-key-that-is-longer-than-32-characters",
        otp_resend_cooldown_seconds=0,
    )


def value[T](records: dict[str, object], key: str, expected_type: type[T]) -> T:
    item = records[key]
    assert isinstance(item, expected_type)
    return item


def _login(
    client: TestClient,
    provider: FakeOtpProvider,
    *,
    phone: str,
    membership_id: str,
) -> str:
    requested = client.post("/api/v1/auth/otp/request", json={"phone": phone})
    assert requested.status_code == 202
    challenge_id = requested.json()["challenge_id"]
    code = provider.deliveries[next(key for key in provider.deliveries if str(key) == challenge_id)]
    verified = client.post(
        "/api/v1/auth/otp/verify",
        json={"challenge_id": challenge_id, "otp": code},
    )
    assert verified.status_code == 200
    memberships = client.post(
        "/api/v1/auth/memberships",
        json={"pre_session_token": verified.json()["pre_session_token"]},
    )
    assert memberships.status_code == 200
    assert any(item["membership_id"] == membership_id for item in memberships.json()["memberships"])
    pre_session = verified.json()["pre_session_token"]
    session = client.post(
        "/api/v1/auth/session",
        json={"pre_session_token": pre_session, "membership_id": membership_id},
    )
    assert session.status_code == 200
    return cast(str, session.json()["access_token"])


def _auth(client: TestClient, token: str) -> None:
    client.headers.update({"Authorization": f"Bearer {token}"})


def _event_payload(
    event_type: str,
    *,
    event_id: str,
    created_at: datetime,
    **extra: object,
) -> dict[str, object]:
    return {
        "client_event_uuid": event_id,
        "event_type": event_type,
        "device_created_at": created_at.isoformat(),
        "installation_identifier": "pilot-device-12",
        "platform": "ANDROID",
        **extra,
    }


def test_owned_tipper_pilot_flow_reconciles_api_and_excel(
    db_session: Session,
    tenant_records: dict[str, object],
) -> None:
    owner_membership = value(tenant_records, "owner_a_membership", CompanyMembership)
    owner_user = value(tenant_records, "owner_a_user", User)
    owner_user.phone_number = "+919876543210"
    db_session.flush()

    provider = FakeOtpProvider()
    storage = InMemoryStorage()
    settings = _settings()
    app = create_app(settings)

    def override_db() -> Iterator[Session]:
        yield db_session

    app.dependency_overrides[session_get_db] = override_db
    app.dependency_overrides[get_otp_provider] = lambda: provider
    app.dependency_overrides[get_object_storage] = lambda: storage
    client = TestClient(app)
    try:
        owner_token = _login(
            client,
            provider,
            phone=owner_user.phone_number,
            membership_id=str(owner_membership.id),
        )
        _auth(client, owner_token)

        site = client.post(
            "/api/v1/admin/sites",
            json={"name": "Pilot Site", "code": "PILOT"},
        )
        tipper = client.post(
            "/api/v1/admin/tippers",
            json={"registration_number": "PILOT-12", "short_name": "Tipper 12"},
        )
        assert site.status_code == 201
        assert tipper.status_code == 201
        site_id = site.json()["id"]
        tipper_id = tipper.json()["id"]

        driver = client.post(
            "/api/v1/admin/people",
            json={"phone": "+919876543221", "display_name": "Pilot Driver", "role": "DRIVER"},
        )
        supervisor = client.post(
            "/api/v1/admin/people",
            json={
                "phone": "+919876543222",
                "display_name": "Pilot Supervisor",
                "role": "SUPERVISOR",
            },
        )
        assert driver.status_code == 201
        assert supervisor.status_code == 201
        driver_membership_id = driver.json()["membership_id"]
        supervisor_membership_id = supervisor.json()["membership_id"]
        access = client.post(
            "/api/v1/admin/supervisor-site-access",
            json={"supervisor_membership_id": supervisor_membership_id, "site_id": site_id},
        )
        assert access.status_code == 201
        assignment = client.post(
            "/api/v1/admin/assignments",
            json={
                "driver_membership_id": driver_membership_id,
                "supervisor_membership_id": supervisor_membership_id,
                "tipper_id": tipper_id,
                "site_id": site_id,
                "starts_at": (datetime.now(UTC) - timedelta(hours=1)).isoformat(),
            },
        )
        assert assignment.status_code == 201

        driver_token = _login(
            client,
            provider,
            phone="+919876543221",
            membership_id=driver_membership_id,
        )
        _auth(client, driver_token)
        created_at = datetime.now(UTC)

        def submit_with_evidence(event_type: str, **extra: object) -> str:
            event_id = str(uuid4())
            upload = client.post(
                "/api/v1/driver/evidence",
                params={"client_event_uuid": event_id},
                files={"file": ("meter.jpg", b"\xff\xd8\xffpilot-image", "image/jpeg")},
            )
            assert upload.status_code == 200
            response = client.post(
                "/api/v1/driver/events",
                json=_event_payload(
                    event_type,
                    event_id=event_id,
                    created_at=created_at,
                    object_reference=upload.json()["object_reference"],
                    **extra,
                ),
            )
            assert response.status_code == 200
            return cast(str, response.json()["event_id"])

        start_event_id = submit_with_evidence(
            "KM_READING",
            reading_type="START_READING",
            reading_value="10000",
        )
        trip_event_ids = []
        for _ in range(8):
            event_id = str(uuid4())
            response = client.post(
                "/api/v1/driver/events",
                json=_event_payload("TRIP_COMPLETE", event_id=event_id, created_at=created_at),
            )
            assert response.status_code == 200
            trip_event_ids.append(response.json()["event_id"])
        diesel_event_id = submit_with_evidence("DIESEL", litres="30")
        end_event_id = submit_with_evidence(
            "KM_READING",
            reading_type="END_READING",
            reading_value="10120",
        )

        supervisor_token = _login(
            client,
            provider,
            phone="+919876543222",
            membership_id=supervisor_membership_id,
        )
        _auth(client, supervisor_token)
        supervisor_events = client.get(
            f"/api/v1/supervisor/sites/{site_id}/events",
            params={"review_date": created_at.date().isoformat()},
        )
        assert supervisor_events.status_code == 200
        event_ids = [item["event_id"] for item in supervisor_events.json()]
        assert set(event_ids) == {start_event_id, *trip_event_ids, diesel_event_id, end_event_id}
        approved = client.post(
            "/api/v1/supervisor/events/verify-batch",
            json={
                "event_ids": event_ids,
                "decision": "APPROVED",
                "expected_status": "PENDING_VERIFICATION",
            },
        )
        assert approved.status_code == 200

        owner_token = _login(
            client,
            provider,
            phone=owner_user.phone_number,
            membership_id=str(owner_membership.id),
        )
        _auth(client, owner_token)
        report_date = created_at.date().isoformat()
        dashboard = client.get(
            "/api/v1/reports/dashboard",
            params={"operational_date": report_date},
        )
        assert dashboard.status_code == 200
        dashboard_json = dashboard.json()
        assert dashboard_json["approved_trip_count"] == 8
        assert float(dashboard_json["total_km"]) == 120
        assert float(dashboard_json["verified_diesel_issued"]) == 30
        site_report = client.get(
            f"/api/v1/reports/sites/{site_id}/daily",
            params={"operational_date": report_date},
        )
        tipper_report = client.get(
            f"/api/v1/reports/tippers/{tipper_id}/daily",
            params={"operational_date": report_date},
        )
        assert site_report.status_code == 200
        assert tipper_report.status_code == 200
        assert site_report.json()["approved_trip_count"] == 8
        assert float(tipper_report.json()[0]["distance_km"]) == 120
        assert float(tipper_report.json()[0]["verified_diesel_issued"]) == 30

        workbook_response = client.get(
            "/api/v1/reports/daily.xlsx",
            params={"operational_date": report_date},
        )
        assert workbook_response.status_code == 200
        workbook = load_workbook(BytesIO(workbook_response.content), read_only=True, data_only=True)
        assert workbook.sheetnames == [
            "Daily Summary",
            "Trip Register",
            "KM Register",
            "Diesel Register",
            "Exceptions",
        ]
        assert workbook["Trip Register"].max_row - 1 == 8
        assert workbook["KM Register"].max_row - 1 == 2
        assert workbook["Diesel Register"].max_row - 1 == 1
        workbook.close()

        closed = client.post(
            f"/api/v1/reports/sites/{site_id}/closure/close",
            params={"operational_date": report_date},
            json={"reason": "Pilot acceptance close"},
        )
        assert closed.status_code == 200
        assert closed.json()["status"] == "CLOSED"
    finally:
        client.close()
