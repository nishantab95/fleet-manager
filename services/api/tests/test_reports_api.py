from __future__ import annotations

from datetime import UTC, date, datetime, timedelta
from io import BytesIO
from typing import cast
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook  # type: ignore[import-untyped]
from sqlalchemy.orm import Session

from fleet_api.db.models import (
    Assignment,
    Company,
    CompanyMembership,
    Site,
    SiteDailyClosureHistory,
    SupervisorSiteAccess,
    Tipper,
)
from fleet_api.domain.enums import SiteStatus
from test_supervisor_api import (
    SupervisorStorage,
    access_token,
    client_for,
    create_driver_client,
    driver_event_payload,
    user_by_name,
    value,
)

pytestmark = pytest.mark.postgres

REPORT_DATE = date(2025, 1, 15)
DAY_START = datetime(2025, 1, 14, 18, tzinfo=UTC)


def add_reporting_assignment(
    db_session: Session,
    records: dict[str, object],
    *,
    site: Site | None = None,
    starts_at: datetime = DAY_START - timedelta(hours=1),
    ends_at: datetime | None = None,
) -> Assignment:
    company = value(records, "company_a", Company)
    assignment = Assignment(
        company_id=company.id,
        driver_membership_id=value(records, "driver_a", CompanyMembership).id,
        supervisor_membership_id=value(records, "supervisor_a", CompanyMembership).id,
        tipper_id=value(records, "tipper_a", Tipper).id,
        site_id=(site or value(records, "site_a", Site)).id,
        starts_at=starts_at,
        ends_at=ends_at,
    )
    db_session.add(assignment)
    db_session.add(
        SupervisorSiteAccess(
            company_id=company.id,
            supervisor_membership_id=assignment.supervisor_membership_id,
            site_id=assignment.site_id,
        )
    )
    db_session.flush()
    return assignment


def owner_client(
    db_session: Session,
    records: dict[str, object],
    *,
    storage: SupervisorStorage | None = None,
) -> TestClient:
    owner = user_by_name(db_session, "Owner A")
    return client_for(
        db_session,
        access_token(db_session, owner, value(records, "owner_a", CompanyMembership)),
        storage=storage,
    )


def supervisor_client(db_session: Session, records: dict[str, object]) -> TestClient:
    supervisor = user_by_name(db_session, "Supervisor A")
    return client_for(
        db_session,
        access_token(
            db_session,
            supervisor,
            value(records, "supervisor_a", CompanyMembership),
        ),
    )


def create_event(
    client: TestClient,
    storage: SupervisorStorage,
    event_type: str,
    created_at: datetime,
    **extra: str,
) -> str:
    client_event_uuid = str(uuid4())
    payload_extra = dict(extra)
    if event_type in {"KM_READING", "DIESEL"}:
        uploaded = client.post(
            "/api/v1/driver/evidence",
            params={"client_event_uuid": client_event_uuid},
            files={"file": ("evidence.jpg", b"\xff\xd8\xffphase6-evidence", "image/jpeg")},
        )
        assert uploaded.status_code == 200
        payload_extra["object_reference"] = uploaded.json()["object_reference"]
    response = client.post(
        "/api/v1/driver/events",
        json=driver_event_payload(
            event_type,
            client_event_uuid=client_event_uuid,
            created_at=created_at,
            **payload_extra,
        ),
    )
    assert response.status_code == 200, response.text
    return cast(str, response.json()["event_id"])


def verify(
    client: TestClient,
    event_id: str,
    decision: str = "APPROVED",
    reason: str | None = None,
) -> None:
    response = client.post(
        f"/api/v1/supervisor/events/{event_id}/verify",
        json={"decision": decision, "reason": reason},
    )
    assert response.status_code == 200, response.text


def test_owner_dashboard_reconciles_site_tipper_excel_and_roles(
    db_session: Session,
    tenant_records: dict[str, object],
) -> None:
    site = value(tenant_records, "site_a", Site)
    site.name = "=Unsafe Site"
    db_session.flush()
    add_reporting_assignment(db_session, tenant_records)
    storage = SupervisorStorage()
    driver = create_driver_client(db_session, tenant_records, storage)
    event_ids: dict[str, list[str]] = {"approved_trips": [], "pending_trips": []}
    try:
        start = create_event(
            driver,
            storage,
            "KM_READING",
            DAY_START + timedelta(hours=1),
            reading_type="START_READING",
            reading_value="10000.00",
        )
        for index in range(8):
            event_ids["approved_trips"].append(
                create_event(
                    driver,
                    storage,
                    "TRIP_COMPLETE",
                    DAY_START + timedelta(hours=2 + index),
                )
            )
        event_ids["pending_trips"].append(
            create_event(driver, storage, "TRIP_COMPLETE", DAY_START + timedelta(hours=12))
        )
        disputed = create_event(driver, storage, "TRIP_COMPLETE", DAY_START + timedelta(hours=13))
        rejected = create_event(driver, storage, "TRIP_COMPLETE", DAY_START + timedelta(hours=14))
        diesel = create_event(
            driver,
            storage,
            "DIESEL",
            DAY_START + timedelta(hours=15),
            litres="30.000",
        )
        create_event(
            driver,
            storage,
            "DIESEL",
            DAY_START + timedelta(hours=16),
            litres="5.000",
        )
        emergency = create_event(
            driver,
            storage,
            "EMERGENCY",
            DAY_START + timedelta(hours=17),
            category="BREAKDOWN",
            description="Open hydraulic warning",
        )
        end = create_event(
            driver,
            storage,
            "KM_READING",
            DAY_START + timedelta(hours=22),
            reading_type="END_READING",
            reading_value="10120.00",
        )
    finally:
        driver.close()

    supervisor = supervisor_client(db_session, tenant_records)
    try:
        verify(supervisor, start)
        verify(supervisor, end)
        for event_id in event_ids["approved_trips"]:
            verify(supervisor, event_id)
        verify(supervisor, diesel)
        verify(supervisor, disputed, "DISPUTED", "Trip details need review")
        verify(supervisor, rejected, "REJECTED", "Rejected by site review")
    finally:
        supervisor.close()

    owner = owner_client(db_session, tenant_records, storage=storage)
    try:
        dashboard = owner.get(
            "/api/v1/reports/dashboard", params={"operational_date": REPORT_DATE.isoformat()}
        )
        assert dashboard.status_code == 200
        dashboard_data = dashboard.json()
        assert dashboard_data["approved_trip_count"] == 8
        assert dashboard_data["pending_trip_count"] == 1
        assert dashboard_data["total_km"] == "120.00"
        assert dashboard_data["verified_diesel_issued"] == "30.000"
        assert dashboard_data["pending_verification_count"] == 2
        assert dashboard_data["unresolved_emergency_count"] == 1
        assert "fuel" not in dashboard.text.lower()

        site_report = owner.get(
            f"/api/v1/reports/sites/{site.id}/daily",
            params={"operational_date": REPORT_DATE.isoformat()},
        )
        assert site_report.status_code == 200
        assert site_report.json()["approved_trip_count"] == dashboard_data["approved_trip_count"]
        tipper_report = owner.get(
            f"/api/v1/reports/tippers/{value(tenant_records, 'tipper_a', Tipper).id}/daily",
            params={"operational_date": REPORT_DATE.isoformat()},
        )
        assert tipper_report.status_code == 200
        tipper_data = tipper_report.json()[0]
        assert tipper_data["approved_trip_count"] == 8
        assert tipper_data["disputed_trip_count"] == 1
        assert tipper_data["rejected_trip_count"] == 1
        assert tipper_data["km_per_approved_trip"] == "15.00"
        assert tipper_data["diesel_issued_per_approved_trip"] == "3.750"
        assert tipper_data["recorded_activity_span_seconds"] == 25_200.0
        assert tipper_data["avg_trip_completion_interval_seconds"] == 3_600.0
        assert tipper_data["median_trip_completion_interval_seconds"] == 3_600.0
        assert tipper_data["longest_trip_gap_seconds"] == 3_600.0
        assert tipper_data["pending_diesel_count"] == 1
        assert tipper_data["disputed_diesel_count"] == 0
        assert tipper_data["closure_status"] == "OPEN"
        foreign_site = value(tenant_records, "site_b", Site)
        assert (
            owner.get(
                f"/api/v1/reports/sites/{foreign_site.id}/daily",
                params={"operational_date": REPORT_DATE.isoformat()},
            ).status_code
            == 404
        )
        evidence = owner.get(f"/api/v1/reports/events/{start}/evidence")
        assert evidence.status_code == 200
        assert evidence.content == b"\xff\xd8\xffphase6-evidence"
        assert evidence.headers["content-type"] == "image/jpeg"
        assert evidence.headers["cache-control"] == "private, no-store"
        assert evidence.headers["x-fleet-evidence-event-type"] == "KM_READING"
        assert evidence.headers["x-fleet-evidence-driver"] == "Driver A"
        exceptions = owner.get(
            "/api/v1/reports/exceptions", params={"operational_date": REPORT_DATE.isoformat()}
        )
        assert exceptions.status_code == 200
        exception_codes = {item["code"] for item in exceptions.json()}
        assert {"TRIP_PENDING", "DIESEL_PENDING", "UNRESOLVED_EMERGENCY"} <= exception_codes

        workbook_response = owner.get(
            "/api/v1/reports/daily.xlsx",
            params={"operational_date": REPORT_DATE.isoformat()},
        )
        assert workbook_response.status_code == 200
        workbook = load_workbook(BytesIO(workbook_response.content), data_only=False)
        assert workbook.sheetnames == [
            "Management Dashboard",
            "Daily Summary",
            "Trip Register",
            "KM Register",
            "Diesel Register",
            "Exceptions",
            "Driver Duty",
        ]
        assert workbook["Driver Duty"]["A1"].value == "Operational Date"
        assert workbook["Daily Summary"]["B2"].value == "'=Unsafe Site"
        assert workbook["Daily Summary"]["J2"].value == 120
        assert workbook["Management Dashboard"]["A8"].value == "Site"
        assert workbook["Management Dashboard"]["L9"].value == 15
        assert workbook["Management Dashboard"]["N9"].value == 3.75
        assert workbook["Management Dashboard"].freeze_panes == "A9"
        trip_rows = list(workbook["Trip Register"].iter_rows(min_row=2, values_only=True))
        assert len(trip_rows) == 11
        assert (
            sum(1 for row in trip_rows if row[5] == "APPROVED")
            == dashboard_data["approved_trip_count"]
        )
        assert workbook["Daily Summary"].freeze_panes == "A2"
        assert workbook["Daily Summary"].auto_filter.ref
        km_rows = list(workbook["KM Register"].iter_rows(min_row=2, values_only=True))
        diesel_rows = list(workbook["Diesel Register"].iter_rows(min_row=2, values_only=True))
        assert any(
            isinstance(row[8], str)
            and row[8].startswith('=HYPERLINK("http://localhost:3000/evidence/')
            and str(start) in row[8]
            and "Open Evidence" in row[8]
            and "Authorization" not in row[8]
            for row in km_rows
        )
        assert any(
            isinstance(row[7], str)
            and row[7].startswith('=HYPERLINK("http://localhost:3000/evidence/')
            and str(diesel) in row[7]
            and "Open Evidence" in row[7]
            and "object" not in row[7].lower()
            for row in diesel_rows
        )
    finally:
        owner.close()

    supervisor_resolver = supervisor_client(db_session, tenant_records)
    try:
        acknowledged = supervisor_resolver.post(
            f"/api/v1/supervisor/events/{emergency}/emergency/acknowledge"
        )
        assert acknowledged.status_code == 200
        resolved = supervisor_resolver.post(
            f"/api/v1/supervisor/events/{emergency}/emergency/resolve"
        )
        assert resolved.status_code == 200
    finally:
        supervisor_resolver.close()

    owner_after_resolution = owner_client(db_session, tenant_records, storage=storage)
    try:
        resolved_dashboard = owner_after_resolution.get(
            "/api/v1/reports/dashboard", params={"operational_date": REPORT_DATE.isoformat()}
        )
        assert resolved_dashboard.status_code == 200
        assert resolved_dashboard.json()["pending_verification_count"] == 2
        assert resolved_dashboard.json()["unresolved_emergency_count"] == 0
    finally:
        owner_after_resolution.close()

    driver_again = create_driver_client(db_session, tenant_records, storage)
    try:
        assert driver_again.get("/api/v1/reports/dashboard").status_code == 403
    finally:
        driver_again.close()
    supervisor_again = supervisor_client(db_session, tenant_records)
    try:
        assert supervisor_again.get("/api/v1/reports/dashboard").status_code == 403
    finally:
        supervisor_again.close()


def test_missing_and_invalid_km_readings_block_closure_with_structured_exceptions(
    db_session: Session,
    tenant_records: dict[str, object],
) -> None:
    site = value(tenant_records, "site_a", Site)
    add_reporting_assignment(db_session, tenant_records)
    storage = SupervisorStorage()
    driver = create_driver_client(db_session, tenant_records, storage)
    try:
        start = create_event(
            driver,
            storage,
            "KM_READING",
            DAY_START + timedelta(hours=1),
            reading_type="START_READING",
            reading_value="100.00",
        )
        invalid_end_uuid = str(uuid4())
        invalid_upload = driver.post(
            "/api/v1/driver/evidence",
            params={"client_event_uuid": invalid_end_uuid},
            files={"file": ("evidence.jpg", b"\xff\xd8\xffphase6-evidence", "image/jpeg")},
        )
        assert invalid_upload.status_code == 200
        invalid_end = driver.post(
            "/api/v1/driver/events",
            json=driver_event_payload(
                "KM_READING",
                client_event_uuid=invalid_end_uuid,
                created_at=DAY_START + timedelta(hours=2),
                reading_type="END_READING",
                reading_value="90.00",
                object_reference=invalid_upload.json()["object_reference"],
            ),
        )
        assert invalid_end.status_code == 422
        assert invalid_end.json()["detail"]["code"] == "INVALID_END_KM"
    finally:
        driver.close()
    supervisor = supervisor_client(db_session, tenant_records)
    try:
        verify(supervisor, start)
    finally:
        supervisor.close()
    owner = owner_client(db_session, tenant_records)
    try:
        report = owner.get(
            f"/api/v1/reports/sites/{site.id}/daily",
            params={"operational_date": REPORT_DATE.isoformat()},
        )
        assert report.status_code == 200
        data = report.json()
        assert data["total_km"] is None
        assert "MISSING_END_READING" in {item["code"] for item in data["tippers"][0]["exceptions"]}
        close = owner.post(
            f"/api/v1/reports/sites/{site.id}/closure/close",
            params={"operational_date": REPORT_DATE.isoformat()},
            json={"reason": None},
        )
        assert close.status_code == 409
        assert close.json()["detail"]["code"] == "CLOSURE_BLOCKED"
        assert any(
            blocker["code"] == "MISSING_END_READING"
            for blocker in close.json()["detail"]["blockers"]
        )
    finally:
        owner.close()


def test_phase6_smoke_fixture_approves_pending_work_then_closes_day(
    db_session: Session,
    tenant_records: dict[str, object],
) -> None:
    site = value(tenant_records, "site_a", Site)
    add_reporting_assignment(db_session, tenant_records)
    storage = SupervisorStorage()
    driver = create_driver_client(db_session, tenant_records, storage)
    try:
        start = create_event(
            driver,
            storage,
            "KM_READING",
            DAY_START + timedelta(hours=1),
            reading_type="START_READING",
            reading_value="10000.00",
        )
        trips = [
            create_event(driver, storage, "TRIP_COMPLETE", DAY_START + timedelta(hours=2 + index))
            for index in range(9)
        ]
        diesel = create_event(
            driver,
            storage,
            "DIESEL",
            DAY_START + timedelta(hours=19),
            litres="30.000",
        )
        end = create_event(
            driver,
            storage,
            "KM_READING",
            DAY_START + timedelta(hours=20),
            reading_type="END_READING",
            reading_value="10120.00",
        )
    finally:
        driver.close()
    supervisor = supervisor_client(db_session, tenant_records)
    try:
        verify(supervisor, start)
        verify(supervisor, end)
        for trip in trips[:8]:
            verify(supervisor, trip)
        verify(supervisor, diesel)
    finally:
        supervisor.close()

    owner = owner_client(db_session, tenant_records, storage=storage)
    try:
        params = {"operational_date": REPORT_DATE.isoformat()}
        before = owner.get("/api/v1/reports/dashboard", params=params)
        assert before.status_code == 200
        assert before.json()["approved_trip_count"] == 8
        assert before.json()["pending_trip_count"] == 1
        assert before.json()["total_km"] == "120.00"
        assert before.json()["verified_diesel_issued"] == "30.000"
        workbook = load_workbook(
            BytesIO(owner.get("/api/v1/reports/daily.xlsx", params=params).content),
            data_only=False,
        )
        assert workbook["Daily Summary"]["F2"].value == 8
        assert workbook["Daily Summary"]["J2"].value == 120
        blocked = owner.post(
            f"/api/v1/reports/sites/{site.id}/closure/close",
            params=params,
            json={"reason": None},
        )
        assert blocked.status_code == 409
        assert any(item["code"] == "TRIP_PENDING" for item in blocked.json()["detail"]["blockers"])
    finally:
        owner.close()

    supervisor = supervisor_client(db_session, tenant_records)
    try:
        verify(supervisor, trips[8])
    finally:
        supervisor.close()
    owner = owner_client(db_session, tenant_records)
    try:
        ready = owner.get(f"/api/v1/reports/sites/{site.id}/closure", params=params)
        assert ready.status_code == 200
        assert ready.json()["status"] == "READY_TO_CLOSE"
        closed = owner.post(
            f"/api/v1/reports/sites/{site.id}/closure/close",
            params=params,
            json={"reason": "Smoke review complete"},
        )
        assert closed.status_code == 200
        assert closed.json()["status"] == "CLOSED"
        after = owner.get("/api/v1/reports/dashboard", params=params)
        assert after.json()["approved_trip_count"] == 9
        assert after.json()["pending_trip_count"] == 0
    finally:
        owner.close()


def test_timezone_boundary_and_assignment_transfer_are_historical_and_not_double_counted(
    db_session: Session,
    tenant_records: dict[str, object],
) -> None:
    company = value(tenant_records, "company_a", Company)
    company.reporting_timezone = "Asia/Kolkata"
    company.operational_day_start_minutes = 360
    first_site = value(tenant_records, "site_a", Site)
    second_site = Site(
        company_id=company.id,
        name="Second Site",
        code="SECOND",
        status=SiteStatus.ACTIVE,
    )
    db_session.add(second_site)
    first_assignment = add_reporting_assignment(
        db_session,
        tenant_records,
        starts_at=datetime(2025, 1, 14, 5, tzinfo=UTC),
        ends_at=datetime(2025, 1, 15, 5, tzinfo=UTC),
    )
    second_assignment = add_reporting_assignment(
        db_session,
        tenant_records,
        site=second_site,
        starts_at=datetime(2025, 1, 15, 5, tzinfo=UTC),
    )
    assert first_assignment.id != second_assignment.id
    db_session.flush()
    storage = SupervisorStorage()
    driver = create_driver_client(db_session, tenant_records, storage)
    try:
        create_event(
            driver,
            storage,
            "KM_READING",
            datetime(2025, 1, 15, 0, tzinfo=UTC),
            reading_type="START_READING",
            reading_value="100.00",
        )
        first_event = create_event(
            driver, storage, "TRIP_COMPLETE", datetime(2025, 1, 15, 0, 45, tzinfo=UTC)
        )
        create_event(
            driver,
            storage,
            "KM_READING",
            datetime(2025, 1, 15, 4, tzinfo=UTC),
            reading_type="END_READING",
            reading_value="105.00",
        )
        create_event(
            driver,
            storage,
            "KM_READING",
            datetime(2025, 1, 15, 5, 15, tzinfo=UTC),
            reading_type="START_READING",
            reading_value="200.00",
        )
        second_event = create_event(
            driver, storage, "TRIP_COMPLETE", datetime(2025, 1, 15, 5, 30, tzinfo=UTC)
        )
    finally:
        driver.close()
    supervisor = supervisor_client(db_session, tenant_records)
    try:
        verify(supervisor, first_event)
        verify(supervisor, second_event)
    finally:
        supervisor.close()
    owner = owner_client(db_session, tenant_records)
    try:
        dashboard = owner.get(
            "/api/v1/reports/dashboard", params={"operational_date": "2025-01-15"}
        )
        assert dashboard.status_code == 200
        data = dashboard.json()
        assert data["operational_date"] == "2025-01-15"
        assert data["workday_start_minutes"] == 360
        assert data["approved_trip_count"] == 2
        first = owner.get(
            f"/api/v1/reports/sites/{first_site.id}/daily",
            params={"operational_date": "2025-01-15"},
        )
        second = owner.get(
            f"/api/v1/reports/sites/{second_site.id}/daily",
            params={"operational_date": "2025-01-15"},
        )
        assert first.json()["approved_trip_count"] == 1
        assert second.json()["approved_trip_count"] == 1
        assert sum(item["approved_trip_count"] for item in data["sites"]) == 2
    finally:
        owner.close()


def test_closure_success_history_reopen_policy_and_company_settings(
    db_session: Session,
    tenant_records: dict[str, object],
) -> None:
    site = value(tenant_records, "site_a", Site)
    add_reporting_assignment(db_session, tenant_records)
    owner = owner_client(db_session, tenant_records)
    try:
        settings = owner.get("/api/v1/admin/company")
        assert settings.status_code == 200
        updated = owner.patch(
            "/api/v1/admin/company",
            json={"reporting_timezone": "UTC", "operational_day_start_minutes": 60},
        )
        assert updated.status_code == 200
        assert updated.json()["operational_day_start_minutes"] == 60
        invalid = owner.patch("/api/v1/admin/company", json={"reporting_timezone": "Not/AZone"})
        assert invalid.status_code == 422
    finally:
        owner.close()

    storage = SupervisorStorage()
    driver = create_driver_client(db_session, tenant_records, storage)
    try:
        start = create_event(
            driver,
            storage,
            "KM_READING",
            datetime(2025, 1, 15, 2, tzinfo=UTC),
            reading_type="START_READING",
            reading_value="100.00",
        )
        end = create_event(
            driver,
            storage,
            "KM_READING",
            datetime(2025, 1, 15, 20, tzinfo=UTC),
            reading_type="END_READING",
            reading_value="120.00",
        )
    finally:
        driver.close()
    supervisor = supervisor_client(db_session, tenant_records)
    try:
        verify(supervisor, start)
        verify(supervisor, end)
    finally:
        supervisor.close()

    supervisor = supervisor_client(db_session, tenant_records)
    try:
        supervisor_close = supervisor.post(
            f"/api/v1/reports/sites/{site.id}/closure/close",
            params={"operational_date": "2025-01-15"},
            json={"reason": "Supervisor daily review complete"},
        )
        assert supervisor_close.status_code == 200, supervisor_close.text
    finally:
        supervisor.close()

    # The changed settings are tested above; use the same clean day for owner reopen.
    owner = owner_client(db_session, tenant_records)
    try:
        already_closed = owner.post(
            f"/api/v1/reports/sites/{site.id}/closure/close",
            params={"operational_date": "2025-01-15"},
            json={"reason": "Daily review complete"},
        )
        assert already_closed.status_code == 409
        missing_reason = owner.post(
            f"/api/v1/reports/sites/{site.id}/closure/reopen",
            params={"operational_date": "2025-01-15"},
            json={"reason": None},
        )
        assert missing_reason.status_code == 422
        reopened = owner.post(
            f"/api/v1/reports/sites/{site.id}/closure/reopen",
            params={"operational_date": "2025-01-15"},
            json={"reason": "Correction required"},
        )
        assert reopened.status_code == 200
        assert reopened.json()["status"] == "READY_TO_CLOSE"
        history = db_session.query(SiteDailyClosureHistory).all()
        assert [item.status.value for item in history] == ["CLOSED", "REOPENED"]
        assert all(item.reason for item in history)
    finally:
        owner.close()
