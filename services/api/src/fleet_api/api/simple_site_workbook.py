from __future__ import annotations

import re
from collections import defaultdict
from dataclasses import dataclass
from datetime import UTC, date, datetime
from decimal import Decimal
from io import BytesIO
from uuid import UUID
from zoneinfo import ZoneInfo

from openpyxl import Workbook  # type: ignore[import-untyped]
from openpyxl.cell.cell import Cell  # type: ignore[import-untyped]
from openpyxl.styles import Alignment, Font, PatternFill  # type: ignore[import-untyped]
from openpyxl.utils import get_column_letter  # type: ignore[import-untyped]
from openpyxl.worksheet.worksheet import Worksheet  # type: ignore[import-untyped]

from fleet_api.db.models import DutySession, FleetAsset
from fleet_api.domain.assets import capabilities_for
from fleet_api.domain.enums import FleetAssetType, OperationalEventType, VerificationStatus
from fleet_api.domain.reporting import AssetDailyReport, ReportMetricState, SitePeriodReport

_TITLE_FILL = PatternFill("solid", fgColor="173C35")
_SECTION_FILL = PatternFill("solid", fgColor="DDEBE7")
_HEADER_FILL = PatternFill("solid", fgColor="2F6F62")
_WHITE_BOLD = Font(bold=True, color="FFFFFF")
_SECTION_FONT = Font(bold=True, color="173C35")
_ILLEGAL_SHEET_CHARACTERS = re.compile(r"[\\/*?:\[\]]")
_ILLEGAL_XML_CHARACTERS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_PENDING_EXCEPTION_CODES = {"KM_PENDING", "TRIP_PENDING", "DIESEL_PENDING"}
_EVENT_STATUS_EXCEPTION_CODES = {"TRIP_DISPUTED"}

FUEL_RATIO_NOTE = (
    "Fuel ratios are based on recorded/verified diesel entries and are not a direct "
    "measurement of actual fuel consumed unless a validated consumption source is available."
)


@dataclass(frozen=True)
class _DailyWorkbookRow:
    operational_date: date
    report: AssetDailyReport
    duty_sessions: tuple[DutySession, ...]


@dataclass(frozen=True)
class _AssetMetrics:
    working_days: int
    approved_trips: int
    total_distance_km: Decimal
    distance_values: int
    distance_incomplete: bool
    total_machine_hours: Decimal
    machine_hour_values: int
    machine_hours_incomplete: bool
    diesel_recorded_l: Decimal
    pending_items: int
    exceptions: int


def _safe_excel_text(value: object) -> object:
    if isinstance(value, str):
        value = _ILLEGAL_XML_CHARACTERS.sub("", value)
        if value.startswith(("=", "+", "-", "@")):
            return "'" + value
    return value


def _set_cell(cell: Cell, value: object) -> None:
    cell.value = _safe_excel_text(value)


def _append_row(worksheet: Worksheet, values: list[object]) -> None:
    worksheet.append([_safe_excel_text(value) for value in values])


def _write_row(worksheet: Worksheet, row_number: int, values: tuple[object, ...]) -> None:
    for column, value in enumerate(values, start=1):
        _set_cell(worksheet.cell(row_number, column), value)


def _asset_display_name(asset: FleetAsset) -> str:
    return asset.short_name or asset.registration_number or asset.asset_code


def _asset_type_label(asset: FleetAsset) -> str:
    return asset.asset_type.value.replace("_", " ")


def _manufacturer_model(asset: FleetAsset) -> str:
    values = [value for value in (asset.manufacturer, asset.model) if value]
    return " / ".join(values) if values else "N/A"


def _site_display_name(report: SitePeriodReport) -> str:
    return report.site.short_name or report.site.name


def _unique_sheet_name(display_name: str, used: set[str]) -> str:
    safe_display_name = _ILLEGAL_XML_CHARACTERS.sub("", display_name)
    raw_name = _ILLEGAL_SHEET_CHARACTERS.sub("_", safe_display_name).strip(" '")
    if raw_name.startswith(("=", "+", "-", "@")):
        raw_name = "_" + raw_name[1:]
    base = raw_name[:31] or "Asset"
    candidate = base
    suffix_number = 2
    while candidate.casefold() in used:
        suffix = f" ({suffix_number})"
        candidate = f"{base[: 31 - len(suffix)]}{suffix}"
        suffix_number += 1
    used.add(candidate.casefold())
    return candidate


def _local_excel_datetime(value: datetime, zone: ZoneInfo) -> datetime:
    if value.tzinfo is None:
        value = value.replace(tzinfo=UTC)
    return value.astimezone(zone).replace(tzinfo=None)


def _pending_count(row: AssetDailyReport) -> int:
    return sum(
        1
        for event in row.events
        if event.event_type != OperationalEventType.EMERGENCY
        and event.verification_status == VerificationStatus.PENDING_VERIFICATION
    )


def _daily_status(row: AssetDailyReport) -> str:
    details: list[str] = []
    for status in (
        VerificationStatus.PENDING_VERIFICATION,
        VerificationStatus.DISPUTED,
        VerificationStatus.REJECTED,
        VerificationStatus.AMENDED,
    ):
        event_types = sorted(
            {
                event.event_type.value
                for event in row.events
                if event.event_type != OperationalEventType.EMERGENCY
                and event.verification_status == status
            }
        )
        if event_types:
            details.append(f"{status.value}: {', '.join(event_types)}")
    details.extend(
        item.code
        for item in row.exceptions
        if item.code not in _PENDING_EXCEPTION_CODES | _EVENT_STATUS_EXCEPTION_CODES
    )
    return " · ".join(dict.fromkeys(details)) if details else row.completeness_status


def _driver_summary(rows: list[_DailyWorkbookRow]) -> str:
    identities = {
        item.report.assignment.driver_membership_id: item.report.driver_name
        for item in rows
        if item.report.driver_name
    }
    if not identities:
        return "N/A"
    if len(identities) == 1:
        return next(iter(identities.values()))
    return "Multiple — see daily records"


def _is_worked_row(item: _DailyWorkbookRow) -> bool:
    return bool(item.duty_sessions) or any(
        event.event_type != OperationalEventType.EMERGENCY for event in item.report.events
    )


def _metrics(rows: list[_DailyWorkbookRow]) -> _AssetMetrics:
    worked_rows = [item for item in rows if _is_worked_row(item)]
    working_dates = {item.operational_date for item in worked_rows}
    distances = [item.report.distance_km for item in rows if item.report.distance_km is not None]
    machine_hours = [
        item.report.machine_hours for item in rows if item.report.machine_hours is not None
    ]
    return _AssetMetrics(
        working_days=len(working_dates),
        approved_trips=sum(item.report.approved_trip_count for item in rows),
        total_distance_km=sum(distances, Decimal("0")),
        distance_values=len(distances),
        distance_incomplete=any(
            item.report.distance_state == ReportMetricState.MISSING for item in worked_rows
        ),
        total_machine_hours=sum(machine_hours, Decimal("0")),
        machine_hour_values=len(machine_hours),
        machine_hours_incomplete=any(
            item.report.machine_hours_state == ReportMetricState.MISSING for item in worked_rows
        ),
        diesel_recorded_l=sum((item.report.verified_diesel_issued for item in rows), Decimal("0")),
        pending_items=sum(_pending_count(item.report) for item in rows),
        exceptions=sum(
            (
                sum(
                    exception.code not in _PENDING_EXCEPTION_CODES | _EVENT_STATUS_EXCEPTION_CODES
                    for exception in item.report.exceptions
                )
                + sum(
                    event.event_type != OperationalEventType.EMERGENCY
                    and event.verification_status == VerificationStatus.DISPUTED
                    for event in item.report.events
                )
            )
            for item in rows
        ),
    )


def _ratio(numerator: Decimal, denominator: Decimal) -> Decimal | str:
    return numerator / denominator if denominator > 0 else "N/A"


def _applicable_metric(
    total: Decimal,
    value_count: int,
    *,
    applicable: bool,
    incomplete: bool = False,
) -> Decimal | int | str:
    if not applicable:
        return "N/A"
    if value_count and not incomplete:
        return total
    return "MISSING"


def _style_title(worksheet: Worksheet, end_column: int) -> None:
    for row_number in (1, 2):
        worksheet.merge_cells(
            start_row=row_number,
            start_column=1,
            end_row=row_number,
            end_column=end_column,
        )
        cell = worksheet.cell(row=row_number, column=1)
        cell.fill = _TITLE_FILL
        cell.font = Font(bold=True, color="FFFFFF", size=16 if row_number == 1 else 13)
        cell.alignment = Alignment(horizontal="center")


def _style_table(
    worksheet: Worksheet,
    *,
    header_row: int,
    last_row: int,
    widths: tuple[float, ...],
) -> None:
    for cell in worksheet[header_row]:
        cell.fill = _HEADER_FILL
        cell.font = _WHITE_BOLD
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    worksheet.auto_filter.ref = (
        f"A{header_row}:{get_column_letter(len(widths))}{max(header_row, last_row)}"
    )
    worksheet.freeze_panes = f"A{header_row + 1}"
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width
    for row in worksheet.iter_rows(min_row=header_row + 1, max_row=last_row):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def _print_setup(worksheet: Worksheet, *, landscape: bool) -> None:
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.page_setup.orientation = "landscape" if landscape else "portrait"
    worksheet.page_margins.left = 0.25
    worksheet.page_margins.right = 0.25
    worksheet.sheet_view.showGridLines = False


def _write_summary(
    worksheet: Worksheet,
    report: SitePeriodReport,
    rows_by_asset: dict[UUID, list[_DailyWorkbookRow]],
    asset_labels: dict[UUID, str],
    *,
    generated_at: datetime,
    zone: ZoneInfo,
) -> None:
    headers = (
        "Asset",
        "Registration",
        "Asset Type",
        "Ownership",
        "Driver / Operator",
        "Days Worked",
        "Approved Trips",
        "Distance KM",
        "Machine Hours",
        "Diesel Recorded L",
        "Pending / Exceptions",
    )
    _set_cell(worksheet["A1"], "FLEET MANAGER")
    _set_cell(worksheet["A2"], "SITE OPERATIONS REPORT")
    _style_title(worksheet, len(headers))
    _set_cell(worksheet["A4"], "Site")
    _set_cell(worksheet["B4"], _site_display_name(report))
    _set_cell(worksheet["A5"], "Report period")
    _set_cell(
        worksheet["B5"],
        f"{report.from_date:%d-%b-%Y} to {report.to_date:%d-%b-%Y}",
    )
    _set_cell(worksheet["A6"], "Generated")
    worksheet["B6"] = _local_excel_datetime(generated_at, zone)
    worksheet["B6"].number_format = "dd-mmm-yyyy hh:mm"
    for cell in (worksheet["A4"], worksheet["A5"], worksheet["A6"]):
        cell.font = _SECTION_FONT

    header_row = 8
    _write_row(worksheet, header_row, headers)
    all_metrics: dict[UUID, _AssetMetrics] = {}
    for asset in report.assets:
        rows = rows_by_asset.get(asset.id, [])
        metrics = _metrics(rows)
        all_metrics[asset.id] = metrics
        capabilities = capabilities_for(asset.asset_type)
        _append_row(
            worksheet,
            [
                asset_labels[asset.id],
                asset.registration_number or "N/A",
                _asset_type_label(asset),
                asset.ownership_type.value,
                _driver_summary(rows),
                metrics.working_days,
                metrics.approved_trips if capabilities.supports_trip_complete else "N/A",
                _applicable_metric(
                    metrics.total_distance_km,
                    metrics.distance_values,
                    applicable=capabilities.supports_odometer,
                    incomplete=metrics.distance_incomplete,
                ),
                _applicable_metric(
                    metrics.total_machine_hours,
                    metrics.machine_hour_values,
                    applicable=capabilities.supports_hour_meter,
                    incomplete=metrics.machine_hours_incomplete,
                ),
                metrics.diesel_recorded_l,
                f"{metrics.pending_items} pending / {metrics.exceptions} exceptions",
            ],
        )
    last_asset_row = header_row + len(report.assets)
    _style_table(
        worksheet,
        header_row=header_row,
        last_row=last_asset_row,
        widths=(24, 18, 18, 13, 27, 14, 16, 16, 16, 19, 30),
    )
    for row in worksheet.iter_rows(min_row=header_row + 1, max_row=last_asset_row):
        for column in (8, 9):
            row[column - 1].number_format = "0.00"
        row[9].number_format = "0.000"

    totals_start = last_asset_row + 3
    worksheet.cell(totals_start, 1, "SITE TOTALS")
    worksheet.cell(totals_start, 1).fill = _SECTION_FILL
    worksheet.cell(totals_start, 1).font = _SECTION_FONT
    tipper_assets = [asset for asset in report.assets if asset.asset_type == FleetAssetType.TIPPER]
    machinery_assets = [
        asset for asset in report.assets if asset.asset_type != FleetAssetType.TIPPER
    ]
    distance_values = sum(all_metrics[asset.id].distance_values for asset in tipper_assets)
    distance_incomplete = any(
        all_metrics[asset.id].distance_incomplete or all_metrics[asset.id].distance_values == 0
        for asset in tipper_assets
    )
    machine_hour_values = sum(
        all_metrics[asset.id].machine_hour_values for asset in machinery_assets
    )
    machine_hours_incomplete = any(
        all_metrics[asset.id].machine_hours_incomplete
        or all_metrics[asset.id].machine_hour_values == 0
        for asset in machinery_assets
    )
    total_rows: list[tuple[str, object]] = [
        ("Total Assets", len(report.assets)),
        ("Total Tippers", len(tipper_assets)),
        ("Total Machinery", len(machinery_assets)),
        (
            "Total Working Days / asset-days",
            sum(item.working_days for item in all_metrics.values()),
        ),
        (
            "Total Approved Trips",
            sum(all_metrics[asset.id].approved_trips for asset in tipper_assets)
            if tipper_assets
            else "N/A",
        ),
        (
            "Total Distance KM",
            sum(
                (all_metrics[asset.id].total_distance_km for asset in tipper_assets),
                Decimal("0"),
            )
            if distance_values and not distance_incomplete
            else "MISSING"
            if tipper_assets
            else "N/A",
        ),
        (
            "Total Machine Hours",
            sum(
                (all_metrics[asset.id].total_machine_hours for asset in machinery_assets),
                Decimal("0"),
            )
            if machine_hour_values and not machine_hours_incomplete
            else "MISSING"
            if machinery_assets
            else "N/A",
        ),
        (
            "Total Diesel Recorded L",
            sum((item.diesel_recorded_l for item in all_metrics.values()), Decimal("0")),
        ),
        ("Pending Items", sum(item.pending_items for item in all_metrics.values())),
        ("Exceptions", sum(item.exceptions for item in all_metrics.values())),
    ]
    for offset, (label, value) in enumerate(total_rows, start=1):
        _set_cell(worksheet.cell(totals_start + offset, 1), label)
        value_cell = worksheet.cell(totals_start + offset, 2)
        _set_cell(value_cell, value)
        if label == "Total Diesel Recorded L":
            value_cell.number_format = "0.000"
        elif label in {"Total Distance KM", "Total Machine Hours"}:
            value_cell.number_format = "0.00"
        else:
            value_cell.number_format = "0"
    worksheet.column_dimensions["A"].width = 34
    worksheet.print_title_rows = f"1:{header_row}"
    _print_setup(worksheet, landscape=True)


def _duty_bounds(
    sessions: tuple[DutySession, ...], zone: ZoneInfo
) -> tuple[datetime | None, datetime | None]:
    if not sessions:
        return None, None
    duty_start = min(session.started_at for session in sessions)
    ended_at = [session.ended_at for session in sessions]
    duty_end = max(value for value in ended_at if value is not None) if all(ended_at) else None
    return (
        _local_excel_datetime(duty_start, zone),
        _local_excel_datetime(duty_end, zone) if duty_end is not None else None,
    )


def _write_asset_sheet(
    worksheet: Worksheet,
    report: SitePeriodReport,
    asset: FleetAsset,
    rows: list[_DailyWorkbookRow],
    asset_label: str,
    *,
    zone: ZoneInfo,
) -> None:
    capabilities = capabilities_for(asset.asset_type)
    metrics = _metrics(rows)
    end_column = 11 if capabilities.supports_odometer else 10
    _set_cell(worksheet["A1"], "ASSET DETAILS")
    _set_cell(worksheet["A2"], asset_label)
    _style_title(worksheet, end_column)
    details: tuple[tuple[str, object], ...] = (
        ("Asset", asset_label),
        ("Registration", asset.registration_number or "N/A"),
        ("Asset Type", _asset_type_label(asset)),
        ("Ownership", asset.ownership_type.value),
        ("Manufacturer / Model", _manufacturer_model(asset)),
        ("Selected Site", _site_display_name(report)),
        ("Driver / Operator", _driver_summary(rows)),
    )
    for row_number, (label, value) in enumerate(details, start=4):
        _set_cell(worksheet.cell(row_number, 1), label)
        _set_cell(worksheet.cell(row_number, 2), value)
        worksheet.cell(row_number, 1).font = _SECTION_FONT

    period_header = 12
    worksheet.cell(period_header, 1, "PERIOD SUMMARY")
    worksheet.cell(period_header, 1).fill = _SECTION_FILL
    worksheet.cell(period_header, 1).font = _SECTION_FONT
    if capabilities.supports_odometer:
        summary_rows: list[tuple[str, object]] = [
            ("Working Days", metrics.working_days),
            ("Approved Trips", metrics.approved_trips),
            (
                "Total Distance KM",
                _applicable_metric(
                    metrics.total_distance_km,
                    metrics.distance_values,
                    applicable=True,
                    incomplete=metrics.distance_incomplete,
                ),
            ),
            ("Diesel Recorded L", metrics.diesel_recorded_l),
            (
                "Average Trips / Working Day",
                Decimal(metrics.approved_trips) / Decimal(metrics.working_days)
                if metrics.working_days
                else "N/A",
            ),
            (
                "Average Distance / Working Day",
                "N/A"
                if not metrics.working_days
                else "MISSING"
                if metrics.distance_incomplete or not metrics.distance_values
                else metrics.total_distance_km / Decimal(metrics.working_days),
            ),
            (
                "Distance per Litre Recorded",
                "N/A"
                if metrics.diesel_recorded_l <= 0
                else "MISSING"
                if metrics.distance_incomplete or not metrics.distance_values
                else _ratio(metrics.total_distance_km, metrics.diesel_recorded_l)
                if metrics.distance_values
                else "MISSING",
            ),
        ]
    else:
        summary_rows = [
            ("Working Days", metrics.working_days),
            (
                "Total Machine Hours",
                _applicable_metric(
                    metrics.total_machine_hours,
                    metrics.machine_hour_values,
                    applicable=True,
                    incomplete=metrics.machine_hours_incomplete,
                ),
            ),
            ("Diesel Recorded L", metrics.diesel_recorded_l),
            (
                "Average Machine Hours / Working Day",
                "N/A"
                if not metrics.working_days
                else "MISSING"
                if metrics.machine_hours_incomplete or not metrics.machine_hour_values
                else metrics.total_machine_hours / Decimal(metrics.working_days),
            ),
            (
                "Litres Recorded / Machine Hour",
                "N/A"
                if metrics.machine_hours_incomplete
                or not metrics.machine_hour_values
                or metrics.total_machine_hours <= 0
                else _ratio(metrics.diesel_recorded_l, metrics.total_machine_hours),
            ),
        ]
    for offset, (label, value) in enumerate(summary_rows, start=1):
        label_cell = worksheet.cell(period_header + offset, 1)
        value_cell = worksheet.cell(period_header + offset, 2)
        _set_cell(label_cell, label)
        _set_cell(value_cell, value)
        label_cell.alignment = Alignment(wrap_text=True, vertical="top")
        if label in {"Working Days", "Approved Trips"}:
            value_cell.number_format = "0"
        elif label == "Diesel Recorded L":
            value_cell.number_format = "0.000"
        else:
            value_cell.number_format = "0.00"

    note_row = period_header + len(summary_rows) + 2
    worksheet.merge_cells(
        start_row=note_row,
        start_column=1,
        end_row=note_row,
        end_column=end_column,
    )
    _set_cell(worksheet.cell(note_row, 1), FUEL_RATIO_NOTE)
    worksheet.cell(note_row, 1).alignment = Alignment(wrap_text=True, vertical="top")
    worksheet.row_dimensions[note_row].height = 36

    daily_title_row = note_row + 2
    worksheet.cell(daily_title_row, 1, "DAILY RECORDS")
    worksheet.cell(daily_title_row, 1).fill = _SECTION_FILL
    worksheet.cell(daily_title_row, 1).font = _SECTION_FONT
    header_row = daily_title_row + 1
    headers: tuple[str, ...]
    widths: tuple[float, ...]
    if capabilities.supports_odometer:
        headers = (
            "Date",
            "Driver",
            "Duty Start",
            "Duty End",
            "Start KM",
            "End KM",
            "Distance KM",
            "Approved Trips",
            "Diesel Recorded L",
            "Distance per Litre Recorded",
            "Pending / Exception Status",
        )
        widths = (34, 24, 13, 13, 14, 14, 15, 16, 18, 23, 42)
    else:
        headers = (
            "Date",
            "Operator",
            "Duty Start",
            "Duty End",
            "Start HMR",
            "End HMR",
            "Machine Hours",
            "Diesel Recorded L",
            "Litres Recorded / Machine Hour",
            "Pending / Exception Status",
        )
        widths = (34, 24, 13, 13, 14, 14, 16, 18, 27, 42)
    _write_row(worksheet, header_row, headers)

    for item in rows:
        duty_start, duty_end = _duty_bounds(item.duty_sessions, zone)
        if capabilities.supports_odometer:
            _append_row(
                worksheet,
                [
                    item.operational_date,
                    item.report.driver_name,
                    duty_start or "",
                    duty_end or "",
                    item.report.start_km if item.report.start_km is not None else "MISSING",
                    item.report.end_km if item.report.end_km is not None else "MISSING",
                    item.report.distance_km if item.report.distance_km is not None else "MISSING",
                    item.report.approved_trip_count,
                    item.report.verified_diesel_issued,
                    (
                        "N/A"
                        if item.report.verified_diesel_issued <= 0
                        else "MISSING"
                        if item.report.distance_km is None
                        else _ratio(
                            item.report.distance_km,
                            item.report.verified_diesel_issued,
                        )
                    ),
                    _daily_status(item.report),
                ],
            )
        else:
            _append_row(
                worksheet,
                [
                    item.operational_date,
                    item.report.driver_name,
                    duty_start or "",
                    duty_end or "",
                    item.report.start_hmr if item.report.start_hmr is not None else "MISSING",
                    item.report.end_hmr if item.report.end_hmr is not None else "MISSING",
                    item.report.machine_hours
                    if item.report.machine_hours is not None
                    else "MISSING",
                    item.report.verified_diesel_issued,
                    (
                        _ratio(
                            item.report.verified_diesel_issued,
                            item.report.machine_hours,
                        )
                        if item.report.machine_hours is not None
                        else "N/A"
                    ),
                    _daily_status(item.report),
                ],
            )

    last_row = header_row + len(rows)
    _style_table(
        worksheet,
        header_row=header_row,
        last_row=last_row,
        widths=widths,
    )
    for row_number in range(header_row + 1, last_row + 1):
        worksheet.cell(row_number, 1).number_format = "dd-mmm-yyyy"
        worksheet.cell(row_number, 3).number_format = "hh:mm"
        worksheet.cell(row_number, 4).number_format = "hh:mm"
        for column in range(5, len(headers)):
            header = headers[column - 1]
            if header == "Diesel Recorded L":
                number_format = "0.000"
            elif header == "Approved Trips":
                number_format = "0"
            else:
                number_format = "0.00"
            worksheet.cell(row_number, column).number_format = number_format
    worksheet.print_title_rows = f"1:{header_row}"
    _print_setup(worksheet, landscape=True)


def build_simple_site_workbook(
    report: SitePeriodReport,
    duty_sessions: list[DutySession],
    *,
    generated_at: datetime | None = None,
) -> bytes:
    """Render one Site/date-range workbook without changing report calculations."""

    zone = ZoneInfo(report.operational_days[0].operational_day.reporting_timezone)
    generated_at = generated_at or datetime.now(UTC)
    sessions_by_assignment_day: dict[tuple[UUID, date], list[DutySession]] = defaultdict(list)
    for session in duty_sessions:
        sessions_by_assignment_day[(session.assignment_id, session.operational_date)].append(
            session
        )

    rows_by_asset: dict[UUID, list[_DailyWorkbookRow]] = defaultdict(list)
    for daily in report.operational_days:
        for row in daily.rows:
            rows_by_asset[row.asset.id].append(
                _DailyWorkbookRow(
                    operational_date=daily.operational_day.operational_date,
                    report=row,
                    duty_sessions=tuple(
                        sessions_by_assignment_day.get(
                            (row.assignment.id, daily.operational_day.operational_date),
                            [],
                        )
                    ),
                )
            )
    for rows in rows_by_asset.values():
        rows.sort(
            key=lambda item: (
                item.operational_date,
                item.report.assignment.starts_at,
                str(item.report.assignment.id),
            )
        )

    display_counts: dict[str, int] = defaultdict(int)
    for asset in report.assets:
        display_counts[_asset_display_name(asset).casefold()] += 1
    asset_labels = {
        asset.id: (
            f"{_asset_display_name(asset)} ({asset.asset_code})"
            if display_counts[_asset_display_name(asset).casefold()] > 1
            else _asset_display_name(asset)
        )
        for asset in report.assets
    }

    workbook = Workbook()
    summary = workbook.active
    summary.title = "SUMMARY"
    _write_summary(
        summary,
        report,
        rows_by_asset,
        asset_labels,
        generated_at=generated_at,
        zone=zone,
    )
    used_sheet_names = {"summary"}
    for asset in report.assets:
        worksheet = workbook.create_sheet(
            _unique_sheet_name(asset_labels[asset.id], used_sheet_names)
        )
        _write_asset_sheet(
            worksheet,
            report,
            asset,
            rows_by_asset.get(asset.id, []),
            asset_labels[asset.id],
            zone=zone,
        )

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
