from __future__ import annotations

from dataclasses import dataclass
from datetime import UTC, date, datetime, timedelta
from io import BytesIO
from typing import NoReturn
from uuid import UUID

from fastapi import APIRouter, Depends, HTTPException, Query, Response, status
from openpyxl import Workbook  # type: ignore[import-untyped]
from openpyxl.cell.cell import MergedCell  # type: ignore[import-untyped]
from openpyxl.styles import Font, PatternFill  # type: ignore[import-untyped]
from openpyxl.utils import get_column_letter  # type: ignore[import-untyped]
from sqlalchemy.orm import Session

from fleet_api.api.dependencies import (
    get_app_settings,
    get_current_session,
    get_object_storage,
    require_owner_admin,
)
from fleet_api.api.schemas import (
    ClosureActionRequest,
    ClosureHistoryResponse,
    ClosureResponse,
    DashboardResponse,
    ReportEventResponse,
    ReportExceptionResponse,
    ReportHistoryResponse,
    SiteDailyReportResponse,
    TipperDailyReportResponse,
)
from fleet_api.auth.service import AuthContext
from fleet_api.core.config import Settings
from fleet_api.db.models import Site
from fleet_api.db.session import get_db
from fleet_api.domain.enums import SiteClosureStatus, VerificationStatus
from fleet_api.domain.errors import (
    ClosureBlockedError,
    ConflictError,
    DomainError,
    NotFoundError,
    ObjectStorageUnavailableError,
    RoleViolationError,
    TenantConsistencyError,
)
from fleet_api.domain.reporting import (
    ClosureSnapshot,
    DashboardReport,
    OperationalDay,
    ReportEvent,
    ReportEvidenceContext,
    ReportException,
    ReportHistory,
    ReportingService,
    SiteDailyReport,
    TipperDailyReport,
)
from fleet_api.storage.objects import ObjectStorage

router = APIRouter(prefix="/api/v1/reports", tags=["reports"])


def _fail(exc: DomainError) -> NoReturn:
    if isinstance(exc, ClosureBlockedError):
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail={"code": "CLOSURE_BLOCKED", "message": str(exc), "blockers": exc.blockers},
        ) from exc
    if isinstance(exc, ObjectStorageUnavailableError):
        http_status, code = status.HTTP_503_SERVICE_UNAVAILABLE, "OBJECT_STORAGE_UNAVAILABLE"
    elif isinstance(exc, NotFoundError):
        http_status, code = status.HTTP_404_NOT_FOUND, "NOT_FOUND"
    elif isinstance(exc, ConflictError):
        http_status, code = status.HTTP_409_CONFLICT, "CONFLICT"
    elif isinstance(exc, TenantConsistencyError | RoleViolationError):
        http_status, code = status.HTTP_403_FORBIDDEN, "FORBIDDEN"
    else:
        http_status, code = status.HTTP_422_UNPROCESSABLE_ENTITY, "VALIDATION_ERROR"
    raise HTTPException(
        status_code=http_status,
        detail={"code": code, "message": str(exc)},
    ) from exc


def _history_response(history: ReportHistory) -> ReportHistoryResponse:
    return ReportHistoryResponse(
        status=history.status,
        actor_name=history.actor_name,
        reason=history.reason,
        created_at=history.created_at,
    )


def _event_response(event: ReportEvent) -> ReportEventResponse:
    return ReportEventResponse(
        event_id=event.event_id,
        event_type=event.event_type,
        assignment_id=event.assignment_id,
        tipper_id=event.tipper_id,
        tipper_registration_number=event.tipper_registration_number,
        site_id=event.site_id,
        site_name=event.site_name,
        driver_name=event.driver_name,
        supervisor_name=event.supervisor_name,
        device_created_at=event.device_created_at,
        server_received_at=event.server_received_at,
        verification_status=event.verification_status,
        reading_type=event.reading_type,
        reading_value=event.reading_value,
        litres=event.litres,
        emergency_category=event.emergency_category,
        emergency_status=event.emergency_status,
        emergency_description=event.emergency_description,
        evidence_available=event.evidence_available,
        verification_history=[_history_response(item) for item in event.verification_history],
    )


def _exception_response(item: ReportException) -> ReportExceptionResponse:
    return ReportExceptionResponse(
        code=item.code,
        description=item.description,
        assignment_id=item.assignment_id,
        tipper_id=item.tipper_id,
        tipper_registration_number=item.tipper_registration_number,
        site_id=item.site_id,
        event_id=item.event_id,
    )


def _duration_seconds(value: timedelta | None) -> float | None:
    return value.total_seconds() if value is not None else None


def _tipper_response(item: TipperDailyReport) -> TipperDailyReportResponse:
    return TipperDailyReportResponse(
        assignment_id=item.assignment.id,
        tipper_id=item.tipper.id,
        registration_number=item.tipper.registration_number,
        short_name=item.tipper.short_name,
        site_id=item.site.id,
        site_name=item.site.name,
        driver_name=item.driver_name,
        supervisor_name=item.supervisor_name,
        assignment_starts_at=item.assignment.starts_at,
        assignment_ends_at=item.assignment.ends_at,
        approved_trip_count=item.approved_trip_count,
        pending_trip_count=item.pending_trip_count,
        disputed_trip_count=item.disputed_trip_count,
        rejected_trip_count=item.rejected_trip_count,
        start_km=item.start_km,
        end_km=item.end_km,
        distance_km=item.distance_km,
        km_per_approved_trip=item.km_per_approved_trip,
        verified_diesel_issued=item.verified_diesel_issued,
        diesel_issued_per_approved_trip=item.diesel_issued_per_approved_trip,
        first_trip_completed_at=item.first_trip_completed_at,
        last_trip_completed_at=item.last_trip_completed_at,
        recorded_activity_span_seconds=_duration_seconds(item.recorded_activity_span),
        avg_trip_completion_interval_seconds=_duration_seconds(item.avg_trip_completion_interval),
        median_trip_completion_interval_seconds=_duration_seconds(
            item.median_trip_completion_interval
        ),
        longest_trip_gap_seconds=_duration_seconds(item.longest_trip_gap),
        pending_diesel_count=item.pending_diesel_count,
        disputed_diesel_count=item.disputed_diesel_count,
        unresolved_emergency_count=item.unresolved_emergency_count,
        missing_start_reading=item.missing_start_reading,
        missing_end_reading=item.missing_end_reading,
        completeness_status=item.completeness_status,
        closure_status=item.closure_status,
        exceptions=[_exception_response(exception) for exception in item.exceptions],
        events=[_event_response(event) for event in item.events],
    )


def _closure_response(site: Site, day: OperationalDay, closure: ClosureSnapshot) -> ClosureResponse:
    return ClosureResponse(
        site_id=site.id,
        site_name=site.name,
        operational_date=day.operational_date,
        reporting_timezone=day.reporting_timezone,
        workday_start_minutes=day.workday_start_minutes,
        status=closure.status,
        blockers=[_exception_response(blocker) for blocker in closure.blockers],
        history=[
            ClosureHistoryResponse(
                status=item.status,
                actor_name=item.actor_name,
                reason=item.reason,
                created_at=item.created_at,
            )
            for item in closure.history
        ],
    )


def _site_response(report: SiteDailyReport) -> SiteDailyReportResponse:
    return SiteDailyReportResponse(
        site_id=report.site.id,
        site_name=report.site.name,
        operational_date=report.operational_day.operational_date,
        reporting_timezone=report.operational_day.reporting_timezone,
        assigned_tippers_count=report.assigned_tippers_count,
        approved_trip_count=report.approved_trip_count,
        pending_trip_count=report.pending_trip_count,
        disputed_trip_count=report.disputed_trip_count,
        total_km=report.total_km,
        verified_diesel_issued=report.verified_diesel_issued,
        missing_reading_count=report.missing_reading_count,
        unresolved_emergency_count=report.unresolved_emergency_count,
        closure=_closure_response(report.site, report.operational_day, report.closure),
        tippers=[_tipper_response(item) for item in report.rows],
    )


def _dashboard_response(report: DashboardReport) -> DashboardResponse:
    return DashboardResponse(
        operational_date=report.operational_day.operational_date,
        reporting_timezone=report.operational_day.reporting_timezone,
        workday_start_minutes=report.operational_day.workday_start_minutes,
        assigned_tippers_count=report.assigned_tippers_count,
        approved_trip_count=report.approved_trip_count,
        pending_trip_count=report.pending_trip_count,
        total_km=report.total_km,
        verified_diesel_issued=report.verified_diesel_issued,
        pending_verification_count=report.pending_verification_count,
        missing_reading_count=report.missing_reading_count,
        unresolved_emergency_count=report.unresolved_emergency_count,
        sites_not_closed_count=report.sites_not_closed_count,
        complete_tippers_count=report.complete_tippers_count,
        sites=[_site_response(site) for site in report.sites],
        exceptions=[_exception_response(item) for item in report.exceptions],
    )


def _service(db: Session, context: AuthContext) -> ReportingService:
    return ReportingService(db, context)


def _evidence_headers(view: ReportEvidenceContext) -> dict[str, str]:
    return {
        "Cache-Control": "private, no-store",
        "Content-Disposition": "inline",
        "X-Fleet-Evidence-Event-Type": view.event.event_type.value,
        "X-Fleet-Evidence-Driver": view.driver_name,
        "X-Fleet-Evidence-Tipper": view.tipper.registration_number,
        "X-Fleet-Evidence-Timestamp": view.event.device_created_at.isoformat(),
    }


@router.get("/dashboard", response_model=DashboardResponse)
def dashboard(
    operational_date: date | None = Query(default=None),
    context: AuthContext = Depends(require_owner_admin),
    db: Session = Depends(get_db),
) -> DashboardResponse:
    try:
        return _dashboard_response(_service(db, context).dashboard(operational_date))
    except DomainError as exc:
        _fail(exc)


@router.get("/sites/{site_id}/daily", response_model=SiteDailyReportResponse)
def site_daily_report(
    site_id: UUID,
    operational_date: date | None = Query(default=None),
    context: AuthContext = Depends(require_owner_admin),
    db: Session = Depends(get_db),
) -> SiteDailyReportResponse:
    try:
        return _site_response(_service(db, context).site_daily(site_id, operational_date))
    except DomainError as exc:
        _fail(exc)


@router.get("/tippers/{tipper_id}/daily", response_model=list[TipperDailyReportResponse])
def tipper_daily_report(
    tipper_id: UUID,
    operational_date: date | None = Query(default=None),
    context: AuthContext = Depends(require_owner_admin),
    db: Session = Depends(get_db),
) -> list[TipperDailyReportResponse]:
    try:
        return [
            _tipper_response(item)
            for item in _service(db, context).tipper_daily(tipper_id, operational_date)
        ]
    except DomainError as exc:
        _fail(exc)


@router.get("/exceptions", response_model=list[ReportExceptionResponse])
def report_exceptions(
    operational_date: date | None = Query(default=None),
    context: AuthContext = Depends(require_owner_admin),
    db: Session = Depends(get_db),
) -> list[ReportExceptionResponse]:
    try:
        return [
            _exception_response(item) for item in _service(db, context).exceptions(operational_date)
        ]
    except DomainError as exc:
        _fail(exc)


@router.get("/sites/{site_id}/closure", response_model=ClosureResponse)
def get_closure(
    site_id: UUID,
    operational_date: date | None = Query(default=None),
    context: AuthContext = Depends(get_current_session),
    db: Session = Depends(get_db),
) -> ClosureResponse:
    try:
        service = _service(db, context)
        day = service.operational_day(operational_date)
        site = service._authorize_site(site_id)
        return _closure_response(site, day, service.closure(site_id, day.operational_date))
    except DomainError as exc:
        _fail(exc)


@router.post("/sites/{site_id}/closure/close", response_model=ClosureResponse)
def close_site(
    site_id: UUID,
    payload: ClosureActionRequest,
    operational_date: date | None = Query(default=None),
    context: AuthContext = Depends(get_current_session),
    db: Session = Depends(get_db),
) -> ClosureResponse:
    try:
        service = _service(db, context)
        closure = service.close_site(site_id, operational_date, reason=payload.reason)
        db.commit()
        day = service.operational_day(operational_date)
        return _closure_response(service._site(site_id), day, closure)
    except DomainError as exc:
        db.rollback()
        _fail(exc)


@router.post("/sites/{site_id}/closure/reopen", response_model=ClosureResponse)
def reopen_site(
    site_id: UUID,
    payload: ClosureActionRequest,
    operational_date: date | None = Query(default=None),
    context: AuthContext = Depends(get_current_session),
    db: Session = Depends(get_db),
) -> ClosureResponse:
    try:
        if not payload.reason:
            raise DomainError("a reopening reason is required")
        service = _service(db, context)
        closure = service.reopen_site(site_id, operational_date, reason=payload.reason)
        db.commit()
        day = service.operational_day(operational_date)
        return _closure_response(service._site(site_id), day, closure)
    except DomainError as exc:
        db.rollback()
        _fail(exc)


def _safe_excel_text(value: object) -> object:
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _event_actor_time(event: ReportEvent) -> tuple[str | None, datetime | None]:
    for item in reversed(event.verification_history):
        if item.status == VerificationStatus.APPROVED:
            return item.actor_name, item.created_at.astimezone(UTC).replace(tzinfo=None)
    return None, None


def _excel_datetime(value: datetime) -> datetime:
    """Excel stores timestamps without timezone metadata; normalize to UTC first."""
    return value.astimezone(UTC).replace(tzinfo=None)


def _excel_local_datetime(value: datetime) -> datetime:
    """Keep a report-timezone timestamp readable in Excel's timezone-naive cells."""
    return value.replace(tzinfo=None)


def _excel_metric(value: object) -> object:
    return value if value is not None else "Unavailable"


def _evidence_application_url(base_url: str, event_id: UUID) -> str:
    return f"{base_url.rstrip('/')}/evidence/{event_id}"


@dataclass(frozen=True)
class _ExcelFormula:
    value: str


def _evidence_formula(base_url: str, event_id: UUID) -> _ExcelFormula:
    url = _evidence_application_url(base_url, event_id)
    return _ExcelFormula(f'=HYPERLINK("{url}","Open Evidence")')


def _build_workbook(report: DashboardReport, *, web_public_base_url: str) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    sheets: dict[str, tuple[list[str], list[list[object]]]] = {}
    summary_headers = [
        "Date",
        "Site",
        "Tipper",
        "Registration",
        "Driver",
        "Approved Trips",
        "Pending Trips",
        "Start KM",
        "End KM",
        "Total KM",
        "Diesel Issued",
        "Emergency Status",
        "Completeness",
        "Closure Status",
        "Rejected Trips",
        "Disputed Trips",
        "Supervisor",
        "KM / Approved Trip",
        "Diesel Issued / Approved Trip",
        "First Trip",
        "Last Trip",
        "Recorded Activity Span",
        "Avg Trip Completion Interval",
        "Median Trip Completion Interval",
        "Longest Trip Gap",
        "Pending Diesel",
        "Disputed Diesel",
        "Missing START",
        "Missing END",
        "Exceptions",
    ]
    management_headers = [
        "Site",
        "Tipper",
        "Driver",
        "Supervisor",
        "Approved Trips",
        "Pending Trips",
        "Rejected Trips",
        "Disputed Trips",
        "Start KM",
        "End KM",
        "Distance KM",
        "KM / Approved Trip",
        "Diesel Issued",
        "Diesel Issued / Approved Trip",
        "First Trip",
        "Last Trip",
        "Recorded Activity Span",
        "Avg Trip Completion Interval",
        "Median Trip Completion Interval",
        "Longest Trip Gap",
        "Pending Diesel",
        "Disputed Diesel",
        "Missing START",
        "Missing END",
        "Unresolved Emergencies",
        "Completeness",
        "Closure",
        "Exceptions",
        "Evidence / Details",
    ]
    summary_rows: list[list[object]] = []
    management_rows: list[list[object]] = []
    trip_rows: list[list[object]] = []
    km_rows: list[list[object]] = []
    diesel_rows: list[list[object]] = []
    exception_rows: list[list[object]] = []
    for site in report.sites:
        for row in site.rows:
            emergency = "UNRESOLVED" if row.unresolved_emergency_count else "CLEAR"
            exception_text = "; ".join(item.code for item in row.exceptions) or "None"
            approved_trip_events = sorted(
                (
                    event
                    for event in row.events
                    if event.event_type.value == "TRIP_COMPLETE"
                    and event.verification_status == VerificationStatus.APPROVED
                ),
                key=lambda event: event.device_created_at,
            )
            previous_approved_by_event: dict[UUID, tuple[datetime, timedelta]] = {}
            previous: ReportEvent | None = None
            for event in approved_trip_events:
                if previous is not None:
                    previous_approved_by_event[event.event_id] = (
                        _excel_datetime(previous.device_created_at),
                        event.device_created_at - previous.device_created_at,
                    )
                previous = event
            first_evidence = next(
                (event for event in row.events if event.evidence_available),
                None,
            )
            evidence_or_details: object = (
                _evidence_formula(web_public_base_url, first_evidence.event_id)
                if first_evidence is not None
                else "No evidence available"
            )
            summary_rows.append(
                [
                    report.operational_day.operational_date,
                    site.site.name,
                    row.tipper.short_name or row.tipper.registration_number,
                    row.tipper.registration_number,
                    row.driver_name,
                    row.approved_trip_count,
                    row.pending_trip_count,
                    row.start_km,
                    row.end_km,
                    row.distance_km,
                    row.verified_diesel_issued,
                    emergency,
                    row.completeness_status,
                    site.closure.status.value,
                    row.rejected_trip_count,
                    row.disputed_trip_count,
                    row.supervisor_name,
                    _excel_metric(row.km_per_approved_trip),
                    _excel_metric(row.diesel_issued_per_approved_trip),
                    _excel_metric(
                        _excel_local_datetime(row.first_trip_completed_at)
                        if row.first_trip_completed_at is not None
                        else None
                    ),
                    _excel_metric(
                        _excel_local_datetime(row.last_trip_completed_at)
                        if row.last_trip_completed_at is not None
                        else None
                    ),
                    _excel_metric(row.recorded_activity_span),
                    _excel_metric(row.avg_trip_completion_interval),
                    _excel_metric(row.median_trip_completion_interval),
                    _excel_metric(row.longest_trip_gap),
                    row.pending_diesel_count,
                    row.disputed_diesel_count,
                    row.missing_start_reading,
                    row.missing_end_reading,
                    exception_text,
                ]
            )
            management_rows.append(
                [
                    site.site.name,
                    row.tipper.short_name or row.tipper.registration_number,
                    row.driver_name,
                    row.supervisor_name,
                    row.approved_trip_count,
                    row.pending_trip_count,
                    row.rejected_trip_count,
                    row.disputed_trip_count,
                    row.start_km,
                    row.end_km,
                    row.distance_km,
                    _excel_metric(row.km_per_approved_trip),
                    row.verified_diesel_issued,
                    _excel_metric(row.diesel_issued_per_approved_trip),
                    _excel_metric(
                        _excel_local_datetime(row.first_trip_completed_at)
                        if row.first_trip_completed_at is not None
                        else None
                    ),
                    _excel_metric(
                        _excel_local_datetime(row.last_trip_completed_at)
                        if row.last_trip_completed_at is not None
                        else None
                    ),
                    _excel_metric(row.recorded_activity_span),
                    _excel_metric(row.avg_trip_completion_interval),
                    _excel_metric(row.median_trip_completion_interval),
                    _excel_metric(row.longest_trip_gap),
                    row.pending_diesel_count,
                    row.disputed_diesel_count,
                    row.missing_start_reading,
                    row.missing_end_reading,
                    row.unresolved_emergency_count,
                    row.completeness_status,
                    site.closure.status.value,
                    exception_text,
                    evidence_or_details,
                ]
            )
            for event in row.events:
                actor, verified_at = _event_actor_time(event)
                if event.event_type.value == "TRIP_COMPLETE":
                    previous_approved = previous_approved_by_event.get(event.event_id)
                    trip_rows.append(
                        [
                            report.operational_day.operational_date,
                            site.site.name,
                            row.tipper.registration_number,
                            row.driver_name,
                            _excel_datetime(event.device_created_at),
                            event.verification_status.value,
                            actor,
                            verified_at,
                            previous_approved[0] if previous_approved else None,
                            previous_approved[1] if previous_approved else None,
                        ]
                    )
                elif event.event_type.value == "KM_READING":
                    km_rows.append(
                        [
                            report.operational_day.operational_date,
                            site.site.name,
                            row.tipper.registration_number,
                            row.driver_name,
                            event.reading_type,
                            event.reading_value,
                            _excel_datetime(event.device_created_at),
                            event.verification_status.value,
                            (
                                _evidence_formula(web_public_base_url, event.event_id)
                                if event.evidence_available
                                else "Not available"
                            ),
                        ]
                    )
                elif event.event_type.value == "DIESEL":
                    diesel_rows.append(
                        [
                            report.operational_day.operational_date,
                            site.site.name,
                            row.tipper.registration_number,
                            row.driver_name,
                            event.litres,
                            _excel_datetime(event.device_created_at),
                            event.verification_status.value,
                            (
                                _evidence_formula(web_public_base_url, event.event_id)
                                if event.evidence_available
                                else "Not available"
                            ),
                        ]
                    )
            for exception in row.exceptions:
                exception_rows.append(
                    [
                        report.operational_day.operational_date,
                        site.site.name,
                        row.tipper.registration_number,
                        exception.code,
                        exception.description,
                        "OPEN",
                    ]
                )
        if site.rows and site.closure.status != SiteClosureStatus.CLOSED:
            exception_rows.append(
                [
                    report.operational_day.operational_date,
                    site.site.name,
                    "",
                    "SITE_NOT_CLOSED",
                    f"Site is {site.closure.status.value}",
                    "OPEN",
                ]
            )
    sheets["Daily Summary"] = (summary_headers, summary_rows)
    sheets["Trip Register"] = (
        [
            "Date",
            "Site",
            "Tipper",
            "Driver",
            "Trip Event Time",
            "Verification Status",
            "Verified By",
            "Verified At",
            "Previous Approved Trip Time",
            "Interval Since Previous Approved Trip",
        ],
        trip_rows,
    )
    sheets["KM Register"] = (
        [
            "Date",
            "Site",
            "Tipper",
            "Driver",
            "Reading Type",
            "KM",
            "Event Time",
            "Verification Status",
            "Evidence",
        ],
        km_rows,
    )
    sheets["Diesel Register"] = (
        [
            "Date",
            "Site",
            "Tipper",
            "Driver",
            "Litres",
            "Event Time",
            "Verification Status",
            "Evidence",
        ],
        diesel_rows,
    )
    sheets["Exceptions"] = (
        ["Date", "Site", "Tipper", "Exception Type", "Description", "Status"],
        exception_rows,
    )
    dashboard = workbook.create_sheet("Management Dashboard")
    dashboard["A1"] = "Management Dashboard"
    dashboard["A1"].font = Font(bold=True, size=16, color="173C35")
    dashboard["A2"] = "Operational Date"
    dashboard["B2"] = report.operational_day.operational_date
    dashboard["D2"] = "Reporting Timezone"
    dashboard["E2"] = report.operational_day.reporting_timezone
    dashboard_summary_headers = [
        "Operational Date",
        "Reporting Timezone",
        "Assigned Tippers",
        "Approved Trips",
        "Total KM",
        "Diesel Issued",
        "Pending Verification",
        "Missing Readings",
        "Unresolved Emergencies",
        "Sites Not Closed",
    ]
    dashboard_summary_values = [
        report.operational_day.operational_date,
        report.operational_day.reporting_timezone,
        report.assigned_tippers_count,
        report.approved_trip_count,
        _excel_metric(report.total_km),
        report.verified_diesel_issued,
        report.pending_verification_count,
        report.missing_reading_count,
        report.unresolved_emergency_count,
        report.sites_not_closed_count,
    ]
    for index, (label, value) in enumerate(
        zip(dashboard_summary_headers, dashboard_summary_values, strict=True), start=1
    ):
        dashboard.cell(row=4, column=index, value=label)
        dashboard.cell(row=5, column=index, value=_safe_excel_text(value))
    dashboard["A7"] = "Tipper Performance"
    dashboard["A7"].font = Font(bold=True, color="173C35")
    group_definitions = [
        (1, 4, "Context", "DDEBF7"),
        (5, 8, "Trip Quality", "E2F0D9"),
        (9, 14, "Distance & Diesel", "FCE4D6"),
        (15, 20, "Trip Timing", "E4DFEC"),
        (21, 29, "Attention & Closure", "FFF2CC"),
    ]
    for start_column, end_column, label, color in group_definitions:
        dashboard.merge_cells(
            start_row=7,
            start_column=start_column,
            end_row=7,
            end_column=end_column,
        )
        cell = dashboard.cell(row=7, column=start_column, value=label)
        cell.font = Font(bold=True, color="173C35")
        cell.fill = PatternFill("solid", fgColor=color)
        for column in range(start_column, end_column + 1):
            dashboard.cell(row=8, column=column).fill = PatternFill("solid", fgColor=color)
    for column, header in enumerate(management_headers, start=1):
        dashboard.cell(row=8, column=column, value=header)
    for row_values in management_rows:
        dashboard.append(
            [
                value.value if isinstance(value, _ExcelFormula) else _safe_excel_text(value)
                for value in row_values
            ]
        )
    dashboard.freeze_panes = "A9"
    dashboard.auto_filter.ref = (
        f"A8:{get_column_letter(len(management_headers))}{max(8, len(management_rows) + 8)}"
    )

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for title, sheet_data in sheets.items():
        headers, sheet_rows = sheet_data
        worksheet = workbook.create_sheet(title)
        worksheet.append(headers)
        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = (
            f"A1:{get_column_letter(len(headers))}{max(1, len(sheet_rows) + 1)}"
        )
        for sheet_row in sheet_rows:
            worksheet.append(
                [
                    value.value if isinstance(value, _ExcelFormula) else _safe_excel_text(value)
                    for value in sheet_row
                ]
            )
        for column in worksheet.columns:
            width = min(max(len(str(cell.value or "")) for cell in column) + 2, 32)
            worksheet.column_dimensions[column[0].column_letter].width = width
        for row in worksheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, timedelta):
                    cell.number_format = '[h]"h "mm"m"'
                elif cell.is_date:
                    cell.number_format = "yyyy-mm-dd hh:mm"
    for row in dashboard.iter_rows():
        for cell in row:
            if isinstance(cell, MergedCell):
                continue
            if isinstance(cell.value, timedelta):
                cell.number_format = '[h]"h "mm"m"'
            elif cell.is_date:
                cell.number_format = "yyyy-mm-dd hh:mm"
    for worksheet, header_row in [(dashboard, 8), (workbook["Daily Summary"], 1)]:
        for column in range(1, worksheet.max_column + 1):
            header = worksheet.cell(row=header_row, column=column).value
            if header in {
                "Start KM",
                "End KM",
                "Distance KM",
                "Total KM",
                "KM / Approved Trip",
                "Diesel Issued",
                "Diesel Issued / Approved Trip",
            }:
                for row_number in range(header_row + 1, worksheet.max_row + 1):
                    worksheet.cell(row=row_number, column=column).number_format = "0.00"
    for column in range(1, dashboard.max_column + 1):
        width = min(
            max(
                len(str(dashboard.cell(row=row_number, column=column).value or ""))
                for row_number in range(1, dashboard.max_row + 1)
            )
            + 2,
            32,
        )
        dashboard.column_dimensions[get_column_letter(column)].width = width
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


@router.get("/daily.xlsx")
def daily_excel(
    operational_date: date | None = Query(default=None),
    context: AuthContext = Depends(require_owner_admin),
    settings: Settings = Depends(get_app_settings),
    db: Session = Depends(get_db),
) -> Response:
    try:
        report = _service(db, context).dashboard(operational_date)
        content = _build_workbook(report, web_public_base_url=settings.web_public_base_url)
        filename = f"fleet-report-{report.operational_day.operational_date.isoformat()}.xlsx"
        return Response(
            content=content,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{filename}"'},
        )
    except DomainError as exc:
        _fail(exc)


@router.get("/events/{event_id}/evidence")
def report_evidence(
    event_id: UUID,
    context: AuthContext = Depends(require_owner_admin),
    storage: ObjectStorage = Depends(get_object_storage),
    db: Session = Depends(get_db),
) -> Response:
    try:
        view = _service(db, context).evidence_context_for_event(event_id)
        content, content_type = storage.read_private(object_key=view.evidence.object_key)
        return Response(
            content=content,
            media_type=content_type,
            headers=_evidence_headers(view),
        )
    except DomainError as exc:
        _fail(exc)
